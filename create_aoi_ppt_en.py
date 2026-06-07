"""
AOI System Customer Acceptance Presentation — English Version
iGuard AOI — Prepreg & CCL Inspection
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ExifTags
import io, os

UP  = "/root/.claude/uploads/7995871e-2904-5c65-bebb-dd5f1424b688"
OUT = "/home/user/gpt-ai-assistant/iGuard_AOI_Acceptance_EN.pptx"

NAVY   = RGBColor(0x0D, 0x2B, 0x55)
BLUE   = RGBColor(0x1A, 0x5F, 0xA8)
CYAN   = RGBColor(0x00, 0xA8, 0xD6)
LGRAY  = RGBColor(0xF4, 0xF6, 0xF9)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DKGRAY = RGBColor(0x33, 0x3C, 0x4A)
ACCENT = RGBColor(0x00, 0xC2, 0x9A)
ORANGE = RGBColor(0xFF, 0x6B, 0x35)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def load_img(fname, max_w=2400, max_h=2400):
    img = Image.open(os.path.join(UP, fname))
    try:
        for k, v in ExifTags.TAGS.items():
            if v == "Orientation":
                ok = k; break
        exif = img._getexif()
        if exif and ok in exif:
            o = exif[ok]
            if o == 3:   img = img.rotate(180, expand=True)
            elif o == 6: img = img.rotate(270, expand=True)
            elif o == 8: img = img.rotate(90,  expand=True)
    except Exception:
        pass
    img.thumbnail((max_w, max_h), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    return buf

def rect(slide, l, t, w, h, fill=None, line=None):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = fill
    else:    s.fill.background()
    if line: s.line.color.rgb = line; s.line.width = Pt(1)
    else:    s.line.fill.background()
    return s

def txt(slide, text, l, t, w, h, size=16, bold=False, color=DKGRAY,
        align=PP_ALIGN.LEFT, font="Arial", italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font

def pic(slide, fname, l, t, w, h):
    slide.shapes.add_picture(load_img(fname), Inches(l), Inches(t), Inches(w), Inches(h))

def header_band(slide, title, sub=None, bg=NAVY, h=1.5):
    rect(slide, 0, 0, 13.33, h, fill=bg)
    rect(slide, 0, h, 13.33, 0.06, fill=CYAN)
    txt(slide, title, 0.45, 0.1, 12.0, 0.85, size=34, bold=True, color=WHITE)
    if sub:
        txt(slide, sub, 0.45, 0.95, 12.0, 0.5, size=17,
            color=RGBColor(0xAA, 0xCC, 0xFF), italic=True)

def footer(slide, n, total=10):
    rect(slide, 0, 7.1, 13.33, 0.4, fill=NAVY)
    txt(slide, "iGuard AOI System  |  Prepreg & CCL Acceptance Review", 0.3, 7.12, 9, 0.35,
        size=12, color=RGBColor(0xAA, 0xBB, 0xCC))
    txt(slide, f"{n} / {total}", 12.0, 7.12, 1.2, 0.35,
        size=12, color=WHITE, align=PP_ALIGN.RIGHT)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 1  Cover
# ══════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
rect(s1, 0, 0, 13.33, 7.5, fill=NAVY)

buf_bg = load_img("6fcd566d-5911.jpg", 1920, 1080)
s1.shapes.add_picture(buf_bg, Inches(0), Inches(0), Inches(13.33), Inches(7.5))

ov = rect(s1, 0, 0, 13.33, 7.5, fill=NAVY)
sf = ov.fill._xPr.find(qn('a:solidFill'))
srgb = sf.find(qn('a:srgbClr'))
al = etree.SubElement(srgb, qn('a:alpha')); al.set('val', '72000')

txt(s1, "iGuard", 0.55, 0.8, 5, 1.0, size=56, bold=True, color=CYAN)
txt(s1, "Automated Optical Inspection System", 0.55, 1.75, 10, 0.65, size=24, color=WHITE)
rect(s1, 0.5, 2.55, 9.0, 0.07, fill=CYAN)
txt(s1, "Prepreg & CCL Material Defect Inspection", 0.55, 2.75, 10, 0.75,
    size=30, bold=True, color=WHITE)
txt(s1, "Customer Acceptance Review", 0.55, 3.6, 10, 0.6,
    size=22, color=RGBColor(0xCC, 0xEE, 0xFF), italic=True)
rect(s1, 0.5, 4.45, 6.0, 0.7, fill=BLUE)
txt(s1, "Acceptance Date: 2026-06-07   |   AOITEK Co., Ltd.", 0.65, 4.52, 5.7, 0.55,
    size=16, color=WHITE)
txt(s1, "1 / 10", 12.0, 7.12, 1.2, 0.35, size=12, color=WHITE, align=PP_ALIGN.RIGHT)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 2  Agenda
# ══════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
rect(s2, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s2, "Acceptance Agenda", "Overview of Today's Review")

agenda = [
    ("01", "System Introduction",      "Architecture, functions, and specifications of iGuard AOI"),
    ("02", "Machine Installation",     "On-site installation, assembly, and setup verification"),
    ("03", "Test Material Confirmation","Prepreg / CCL roll sample specification review"),
    ("04", "Material Loading & Feed",  "Roll loading, tension control, and feed path verification"),
    ("05", "System Operation Test",    "Full-speed run, AOI scanning, and real-time processing"),
    ("06", "Defect Detection Demo",    "Live defect images, classification, and grading results"),
    ("07", "Software Interface Review","HMI operations, reporting, and data export functions"),
    ("08", "Customer Discussion",      "Customer feedback, Q&A, and engineer responses"),
    ("09", "Acceptance Conclusion",    "Passed items / Action items / Next steps"),
]

for i, (num, title, desc) in enumerate(agenda):
    col = i % 2; row = i // 2
    x = 0.4 + col * 6.5; y = 1.7 + row * 1.12
    rect(s2, x, y, 6.1, 0.98, fill=WHITE, line=BLUE)
    rect(s2, x, y, 0.72, 0.98, fill=BLUE)
    txt(s2, num, x+0.02, y+0.17, 0.68, 0.6, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s2, title, x+0.82, y+0.06, 5.1, 0.45, size=17, bold=True, color=NAVY)
    txt(s2, desc,  x+0.82, y+0.52, 5.1, 0.42, size=13, color=DKGRAY, italic=True)

footer(s2, 2)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 3  System Overview
# ══════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
rect(s3, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s3, "iGuard AOI — System Overview", "Architecture & Capabilities")

pic(s3, "6fcd566d-5911.jpg", 0.35, 1.7, 5.4, 5.2)

specs = [
    (BLUE,   "Inspection Target",  "Prepreg (glass-fiber prepreg rolls) / CCL (copper-clad laminate)"),
    (ACCENT, "Defect Types",       "Foreign matter, scratch, dent, void, crease, bubble, edge defect"),
    (NAVY,   "Scanning Method",    "Full-width line-scan camera with real-time row-by-row comparison"),
    (ORANGE, "Output Functions",   "Live marking, defect grading, report export, AI-assisted review"),
    (BLUE,   "System Brand",       "iGuard  ·  AOITEK Co., Ltd."),
]
for i, (c, label, val) in enumerate(specs):
    y = 1.75 + i * 1.08
    rect(s3, 6.1, y, 6.8, 0.95, fill=WHITE, line=c)
    rect(s3, 6.1, y, 1.9, 0.95, fill=c)
    txt(s3, label, 6.12, y+0.2, 1.86, 0.55, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s3, val,   8.1,  y+0.12, 4.7, 0.72, size=14, color=DKGRAY)

footer(s3, 3)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 4  Machine Installation
# ══════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(BLANK)
rect(s4, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s4, "Machine Installation & Setup", "Equipment Assembly Verification")

photos4 = [
    ("c134e148-5922.jpg", "Engineer inspecting machine frame structure"),
    ("36904af0-5920.jpg", "Machine framework visual inspection"),
    ("f70770e5-5919.jpg", "Team collaborative assembly and calibration"),
    ("566f775c-5918.jpg", "Multi-person on-site assembly verification"),
]
for i, (fn, cap) in enumerate(photos4):
    col = i % 2; row = i // 2
    x = 0.35 + col * 6.5; y = 1.65 + row * 2.75
    pic(s4, fn, x, y, 6.0, 2.45)
    rect(s4, x, y+2.45, 6.0, 0.28, fill=NAVY)
    txt(s4, cap, x+0.1, y+2.47, 5.8, 0.26, size=13, color=WHITE, align=PP_ALIGN.CENTER)

footer(s4, 4)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 5  Test Material
# ══════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(BLANK)
rect(s5, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s5, "Test Material Confirmation", "Prepreg / CCL Roll Sample Review")

pic(s5, "61bc114c-5915.jpg", 0.35, 1.65, 5.9, 2.7)
rect(s5, 0.35, 4.35, 5.9, 0.28, fill=NAVY)
txt(s5, "Prepreg roll (olive-green, glass-fiber epoxy prepreg)",
    0.45, 4.37, 5.7, 0.26, size=13, color=WHITE, align=PP_ALIGN.CENTER)

pic(s5, "84500fd0-5908.jpg", 0.35, 4.72, 5.9, 2.22)
rect(s5, 0.35, 6.94, 5.9, 0.28, fill=NAVY)
txt(s5, "Prepreg surface detail (visible crease / minor defect)",
    0.45, 6.96, 5.7, 0.26, size=13, color=WHITE, align=PP_ALIGN.CENTER)

rect(s5, 6.55, 1.65, 6.45, 5.55, fill=WHITE, line=BLUE)
txt(s5, "Material Specifications", 6.75, 1.75, 6.1, 0.55, size=20, bold=True, color=NAVY)
rect(s5, 6.55, 2.3, 6.45, 0.06, fill=CYAN)

mat_info = [
    ("Material Type", "Prepreg (glass-fiber epoxy prepreg) / CCL (copper-clad laminate)"),
    ("Supplier",      "DOOSAN Corporation"),
    ("Appearance",    "Roll form, olive-green / brown, width approx. 600 mm"),
    ("Key Defects",   "Crease, surface scratch, foreign matter, dent, bubble"),
    ("Roll Dimension","Per production spec (including core weight)"),
    ("Loading Method","Unwind spindle with automatic tension controller"),
]
for i, (k, v) in enumerate(mat_info):
    y = 2.5 + i * 0.75
    rect(s5, 6.6, y, 1.85, 0.62, fill=BLUE)
    txt(s5, k, 6.65, y+0.1, 1.75, 0.45, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s5, v, 8.55, y+0.1, 4.3, 0.55, size=13, color=DKGRAY)

footer(s5, 5)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 6  Material Loading & Feed
# ══════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(BLANK)
rect(s6, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s6, "Material Loading & Feed Verification", "Roll Threading and Feed Path Confirmation")

photos6 = [
    ("db7d8cf6-5914.jpg", "Engineer manually mounting the roll"),
    ("d4329ee6-5916.jpg", "Threading material through the feed rollers"),
    ("30c03f19-5917.jpg", "Full machine feed path confirmed"),
]
for i, (fn, cap) in enumerate(photos6):
    x = 0.3 + i * 4.35
    pic(s6, fn, x, 1.65, 4.1, 4.9)
    rect(s6, x, 6.55, 4.1, 0.28, fill=NAVY)
    txt(s6, cap, x+0.1, 6.57, 3.9, 0.26, size=13, color=WHITE, align=PP_ALIGN.CENTER)

steps = ["① Confirm Roll", "② Mount Spindle", "③ Thread Material", "④ Set Tension", "⑤ Trial Run"]
for i, s in enumerate(steps):
    c = [BLUE, CYAN, ACCENT, ORANGE, NAVY][i]
    x = 0.3 + i * 2.6
    rect(s6, x, 6.95, 2.45, 0.4, fill=c)
    txt(s6, s, x+0.05, 6.97, 2.35, 0.36,
        size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

footer(s6, 6)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 7  System Operation & Scanning
# ══════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(BLANK)
rect(s7, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s7, "System Operation & Scanning Test", "Functional Verification Checklist")

pic(s7, "e6a2d99f-5909.jpg",  0.35, 1.65, 5.9, 2.55)
rect(s7, 0.35, 4.2, 5.9, 0.28, fill=NAVY)
txt(s7, "AOI scanning unit in operation", 0.45, 4.22, 5.7, 0.26,
    size=13, color=WHITE, align=PP_ALIGN.CENTER)

pic(s7, "ee17691c-5906.jpg", 0.35, 4.58, 5.9, 2.25)
rect(s7, 0.35, 6.83, 5.9, 0.28, fill=NAVY)
txt(s7, "Engineer verifying scan status on-site", 0.45, 6.85, 5.7, 0.26,
    size=13, color=WHITE, align=PP_ALIGN.CENTER)

rect(s7, 6.55, 1.65, 6.45, 5.55, fill=WHITE, line=ACCENT)
txt(s7, "Verification Checklist", 6.75, 1.75, 6.1, 0.55, size=20, bold=True, color=NAVY)
rect(s7, 6.55, 2.3, 6.45, 0.05, fill=ACCENT)

checks = [
    "Machine startup normal — no fault alarms",
    "Feed speed stable — no tension fluctuation",
    "Line-scan camera imaging clear and sharp",
    "Real-time display smooth (≥ 30 fps)",
    "AOI algorithm loaded and running correctly",
    "Emergency stop (E-Stop) function verified",
    "Control cabinet wiring neat — no loose connections",
]
for i, desc in enumerate(checks):
    y = 2.45 + i * 0.72
    rect(s7, 6.6, y, 0.55, 0.58, fill=ACCENT)
    txt(s7, "✔", 6.6, y+0.08, 0.55, 0.45,
        size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s7, desc, 7.25, y+0.1, 5.6, 0.48, size=14, color=DKGRAY)

footer(s7, 7)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 8  Defect Detection Results
# ══════════════════════════════════════════════════════════════════════════
s8 = prs.slides.add_slide(BLANK)
rect(s8, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s8, "Defect Detection Results", "Live AOI Output — Defect Classification & Grading")

pic(s8, "6dd1d6e3-5910.jpg", 0.35, 1.65, 6.1, 5.55)
rect(s8, 0.35, 7.2, 6.1, 0.28, fill=NAVY)
txt(s8, "AOI software displaying real-time CCL scan results with defect marking",
    0.45, 7.22, 5.9, 0.26, size=12, color=WHITE, align=PP_ALIGN.CENTER)

rect(s8, 6.75, 1.65, 6.25, 0.65, fill=BLUE)
txt(s8, "Defect Classification Results", 6.85, 1.73, 6.05, 0.52,
    size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

defects = [
    ("Crease",          "Grade B", ORANGE, "Detected"),
    ("Surface Scratch", "Grade A", BLUE,   "Detected"),
    ("Dent",            "Grade B", ORANGE, "Detected"),
    ("Foreign Matter",  "Grade A", BLUE,   "Detected"),
    ("Edge Defect",     "Grade C", ACCENT, "Detected"),
    ("Normal",          "PASS",    ACCENT, "Passed"),
]
for i, (name, grade, c, status) in enumerate(defects):
    y = 2.38 + i * 0.78
    bg = LGRAY if i % 2 == 0 else WHITE
    rect(s8, 6.75, y, 6.25, 0.75, fill=bg)
    txt(s8, name, 6.85, y+0.12, 3.2, 0.52, size=14, color=DKGRAY)
    rect(s8, 10.15, y+0.1, 1.35, 0.52, fill=c)
    txt(s8, grade, 10.15, y+0.12, 1.35, 0.5,
        size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s8, status, 11.6, y+0.12, 1.3, 0.52,
        size=13, color=DKGRAY, align=PP_ALIGN.CENTER)

footer(s8, 8)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 9  Customer Review & Discussion
# ══════════════════════════════════════════════════════════════════════════
s9 = prs.slides.add_slide(BLANK)
rect(s9, 0, 0, 13.33, 7.5, fill=LGRAY)
header_band(s9, "Customer Review & On-site Discussion", "Joint Inspection Walkthrough")

photos9 = [
    ("9d6dd958-5907.jpg", "Customer and engineers conduct joint machine review"),
    ("c8d31686-5904.jpg", "On-site discussion of machine configuration"),
    ("6f430e2d-5901.jpg", "Customer representative examines mechanical details"),
    ("8a891a71-5899.jpg", "Sensor and electrical wiring verification"),
]
for i, (fn, cap) in enumerate(photos9):
    col = i % 2; row = i // 2
    x = 0.35 + col * 6.5; y = 1.65 + row * 2.75
    pic(s9, fn, x, y, 6.0, 2.45)
    rect(s9, x, y+2.45, 6.0, 0.28, fill=NAVY)
    txt(s9, cap, x+0.1, y+2.47, 5.8, 0.26, size=13, color=WHITE, align=PP_ALIGN.CENTER)

footer(s9, 9)

# ══════════════════════════════════════════════════════════════════════════
#  SLIDE 10  Acceptance Conclusion
# ══════════════════════════════════════════════════════════════════════════
s10 = prs.slides.add_slide(BLANK)
rect(s10, 0, 0, 13.33, 7.5, fill=NAVY)
rect(s10, 0, 0,    13.33, 0.12, fill=CYAN)
rect(s10, 0, 7.38, 13.33, 0.12, fill=CYAN)

txt(s10, "Acceptance Conclusion", 0.5, 0.2, 12, 0.9,
    size=38, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s10, "iGuard AOI System  ·  Prepreg & CCL Inspection  ·  2026-06-07", 0.5, 1.05, 12, 0.55,
    size=17, color=CYAN, align=PP_ALIGN.CENTER, italic=True)

rect(s10, 0.4, 1.8, 5.9, 0.55, fill=ACCENT)
txt(s10, "✔  Items Accepted", 0.55, 1.85, 5.6, 0.48, size=18, bold=True, color=WHITE)
pass_items = [
    "Machine structure assembled correctly, no loose parts",
    "Roll loading and material feed path smooth",
    "Line-scan camera image quality confirmed",
    "Real-time defect detection functioning properly",
    "HMI software interface clear and operable",
    "Emergency stop (E-Stop) function verified and passed",
]
for i, item in enumerate(pass_items):
    txt(s10, f"  ●  {item}", 0.5, 2.48 + i*0.68, 5.7, 0.6,
        size=15, color=RGBColor(0xCC, 0xFF, 0xEE))

rect(s10, 6.85, 1.8, 6.1, 0.55, fill=ORANGE)
txt(s10, "⚡  Action Items", 7.0, 1.85, 5.8, 0.48, size=18, bold=True, color=WHITE)
action_items = [
    "Sign and finalize the acceptance report",
    "Complete machine cleaning and packaging before shipment",
    "Provide operation training and maintenance manual",
    "Confirm on-site installation schedule",
    "Finalize after-sales warranty terms in writing",
]
for i, item in enumerate(action_items):
    txt(s10, f"  →  {item}", 6.95, 2.48 + i*0.68, 5.8, 0.6,
        size=15, color=RGBColor(0xFF, 0xE0, 0xCC))

txt(s10, "Thank you for your participation  ·  AOITEK Co., Ltd.  ·  2026-06-07",
    0, 7.05, 13.33, 0.38, size=14,
    color=RGBColor(0xAA, 0xCC, 0xFF), align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Saved → {OUT}")
